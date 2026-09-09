import io
import math
import uuid
from datetime import date

from fastapi import APIRouter, File, HTTPException, Query, Request, UploadFile, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select

from app.core.deps import CurrentUser, DbSession, get_client_ip
from app.models.entities import (
    Contract,
    Customer,
    CustomerSession,
    CustomerSubscription,
    CustomerStatus,
    Document,
    HoursTransaction,
    HoursTransactionType,
    Package,
    Payment,
    RoomBooking,
    SubscriptionStatus,
)
from app.schemas.customers import (
    CustomerCreate,
    CustomerProfileResponse,
    CustomerResponse,
    CustomerUpdate,
    HoursAdjustmentRequest,
)
from app.services.audit import log_audit
from app.services.hours import add_hours_transaction, get_hours_summary

router = APIRouter(prefix="/customers", tags=["العملاء"])


def _next_customer_code(db) -> str:
    count = db.scalar(select(func.count()).select_from(Customer)) or 0
    return f"C{date.today().year}{str(count + 1).zfill(5)}"


@router.get("")
def list_customers(
    db: DbSession,
    user: CurrentUser,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=100),
    search: str | None = None,
    status: str | None = None,
):
    q = select(Customer).where(Customer.deleted_at.is_(None))
    if search:
        term = f"%{search}%"
        q = q.where(
            or_(
                Customer.full_name.ilike(term),
                Customer.phone.ilike(term),
                Customer.national_id.ilike(term),
                Customer.company_name.ilike(term),
                Customer.customer_code.ilike(term),
            )
        )
    if status:
        q = q.where(Customer.status == status)
    total = db.scalar(select(func.count()).select_from(q.subquery())) or 0
    items = db.scalars(
        q.order_by(Customer.created_at.desc()).offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {
        "items": [CustomerResponse.model_validate(i) for i in items],
        "total": total,
        "page": page,
        "page_size": page_size,
        "pages": math.ceil(total / page_size) if total else 0,
    }


@router.get("/import-template")
def download_import_template(user: CurrentUser):
    wb = Workbook()
    ws = wb.active
    ws.title = "customers"
    ws.append(["full_name", "national_id", "phone", "email", "company_name", "address"])
    ws.append(["أحمد محمد", "29001011234567", "01012345678", "email@example.com", "شركة مثال", "6 أكتوبر"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=customers_import_template.xlsx"},
    )


@router.post("/import")
async def import_customers(request: Request, db: DbSession, user: CurrentUser, file: UploadFile = File(...)):
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xls")):
        raise HTTPException(400, "يجب رفع ملف Excel (.xlsx)")
    content = await file.read()
    wb = load_workbook(io.BytesIO(content), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    created = 0
    errors: list[str] = []
    for idx, row in enumerate(rows, start=2):
        if not row or not row[0]:
            continue
        try:
            data = CustomerCreate(
                full_name=str(row[0]).strip(),
                national_id=str(row[1]).strip(),
                phone=str(row[2]).strip(),
                email=str(row[3]).strip() if row[3] else None,
                company_name=str(row[4]).strip() if row[4] else None,
                address=str(row[5]).strip() if row[5] else None,
            )
            if db.scalar(select(Customer).where(Customer.national_id == data.national_id, Customer.deleted_at.is_(None))):
                errors.append(f"سطر {idx}: الرقم القومي مسجل")
                continue
            if db.scalar(select(Customer).where(Customer.phone == data.phone, Customer.deleted_at.is_(None))):
                errors.append(f"سطر {idx}: الهاتف مسجل")
                continue
            customer = Customer(customer_code=_next_customer_code(db), **data.model_dump())
            db.add(customer)
            created += 1
        except Exception as exc:
            errors.append(f"سطر {idx}: {exc}")
    log_audit(db, user_id=user.id, action="import", module="customers",
              new_value={"created": created, "errors": len(errors)}, ip_address=get_client_ip(request))
    db.commit()
    return {"created": created, "errors": errors}


@router.post("", response_model=CustomerResponse, status_code=201)
def create_customer(data: CustomerCreate, request: Request, db: DbSession, user: CurrentUser):
    if db.scalar(select(Customer).where(Customer.national_id == data.national_id, Customer.deleted_at.is_(None))):
        raise HTTPException(400, "الرقم القومي مسجل مسبقًا")
    if db.scalar(select(Customer).where(Customer.phone == data.phone, Customer.deleted_at.is_(None))):
        raise HTTPException(400, "رقم الهاتف مسجل مسبقًا")
    customer = Customer(customer_code=_next_customer_code(db), **data.model_dump())
    db.add(customer)
    db.flush()
    log_audit(db, user_id=user.id, action="create", module="customers", record_id=str(customer.id),
              new_value=data.model_dump(), ip_address=get_client_ip(request))
    db.commit()
    db.refresh(customer)
    return customer


@router.delete("/hours-transactions/{transaction_id}")
def delete_hours_transaction(transaction_id: uuid.UUID, request: Request, db: DbSession, user: CurrentUser):
    tx = db.get(HoursTransaction, transaction_id)
    if not tx:
        raise HTTPException(404, "العملية غير موجودة")
    linked = db.scalars(
        select(CustomerSession).where(CustomerSession.hours_transaction_id == transaction_id)
    ).all()
    for session in linked:
        session.hours_transaction_id = None
        session.hours_deducted = None
    log_audit(
        db,
        user_id=user.id,
        action="delete",
        module="hours",
        record_id=str(transaction_id),
        old_value={"amount": float(tx.amount), "customer_id": str(tx.customer_id)},
        ip_address=get_client_ip(request),
    )
    db.delete(tx)
    db.commit()
    return {"message": "تم حذف العملية"}


@router.get("/{customer_id}", response_model=CustomerProfileResponse)
def get_customer(customer_id: uuid.UUID, db: DbSession, user: CurrentUser):
    customer = db.get(Customer, customer_id)
    if not customer or customer.deleted_at:
        raise HTTPException(404, "العميل غير موجود")
    sub = db.scalar(
        select(CustomerSubscription)
        .where(CustomerSubscription.customer_id == customer_id, CustomerSubscription.status == SubscriptionStatus.ACTIVE.value)
        .order_by(CustomerSubscription.created_at.desc())
    )
    active_sub = None
    if sub:
        pkg = db.get(Package, sub.package_id)
        active_sub = {
            "id": str(sub.id),
            "package_name": pkg.name if pkg else None,
            "subscription_type": sub.subscription_type,
            "status": sub.status,
            "start_date": str(sub.start_date),
            "end_date": str(sub.end_date) if sub.end_date else None,
        }
    return CustomerProfileResponse(
        **CustomerResponse.model_validate(customer).model_dump(),
        hours_summary=get_hours_summary(db, customer_id),
        active_subscription=active_sub,
        contracts_count=db.scalar(select(func.count()).where(Contract.customer_id == customer_id, Contract.deleted_at.is_(None))) or 0,
        documents_count=db.scalar(select(func.count()).where(Document.customer_id == customer_id, Document.deleted_at.is_(None))) or 0,
        bookings_count=db.scalar(select(func.count()).where(RoomBooking.customer_id == customer_id, RoomBooking.deleted_at.is_(None))) or 0,
        payments_count=db.scalar(select(func.count()).where(Payment.customer_id == customer_id)) or 0,
    )


@router.patch("/{customer_id}", response_model=CustomerResponse)
def update_customer(customer_id: uuid.UUID, data: CustomerUpdate, request: Request, db: DbSession, user: CurrentUser):
    customer = db.get(Customer, customer_id)
    if not customer or customer.deleted_at:
        raise HTTPException(404, "العميل غير موجود")
    old = CustomerResponse.model_validate(customer).model_dump()
    updates = data.model_dump(exclude_unset=True)
    if "phone" in updates:
        dup = db.scalar(select(Customer).where(Customer.phone == updates["phone"], Customer.id != customer_id, Customer.deleted_at.is_(None)))
        if dup:
            raise HTTPException(400, "رقم الهاتف مسجل مسبقًا")
    for k, v in updates.items():
        setattr(customer, k, v)
    log_audit(db, user_id=user.id, action="update", module="customers", record_id=str(customer_id),
              old_value=old, new_value=updates, ip_address=get_client_ip(request))
    db.commit()
    db.refresh(customer)
    return customer


@router.delete("/{customer_id}")
def delete_customer(customer_id: uuid.UUID, request: Request, db: DbSession, user: CurrentUser):
    customer = db.get(Customer, customer_id)
    if not customer or customer.deleted_at:
        raise HTTPException(404, "العميل غير موجود")
    from datetime import datetime, timezone
    customer.deleted_at = datetime.now(timezone.utc)
    customer.status = CustomerStatus.INACTIVE.value
    log_audit(db, user_id=user.id, action="delete", module="customers", record_id=str(customer_id),
              ip_address=get_client_ip(request))
    db.commit()
    return {"message": "تم حذف العميل"}


@router.post("/{customer_id}/hours")
def adjust_hours(customer_id: uuid.UUID, data: HoursAdjustmentRequest, request: Request, db: DbSession, user: CurrentUser):
    customer = db.get(Customer, customer_id)
    if not customer or customer.deleted_at:
        raise HTTPException(404, "العميل غير موجود")
    from decimal import Decimal
    tx_type = data.transaction_type
    if tx_type not in ("bonus", "adjustment", "correction"):
        tx_type = HoursTransactionType.BONUS.value
    amount = Decimal(str(data.amount))
    if tx_type == HoursTransactionType.USAGE.value and amount > 0:
        amount = -amount
    try:
        tx = add_hours_transaction(
            db, customer_id=customer_id, amount=amount,
            transaction_type=tx_type, reason=data.reason, created_by=user.id,
            allow_negative=user.is_superuser,
        )
    except ValueError as e:
        raise HTTPException(400, str(e))
    log_audit(db, user_id=user.id, action="hours_adjustment", module="hours",
              record_id=str(customer_id), new_value={"amount": float(amount), "reason": data.reason},
              ip_address=get_client_ip(request))
    db.commit()
    return {"message": "تم تسجيل العملية", "balance": float(tx.balance_after or 0)}


@router.get("/{customer_id}/hours")
def customer_hours(customer_id: uuid.UUID, db: DbSession, user: CurrentUser):
    return get_hours_summary(db, customer_id)
