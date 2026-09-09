import io
import uuid
from datetime import date

from fastapi import APIRouter, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.deps import CurrentUser, DbSession
from app.models.entities import Customer, ImportJob, ImportRow, Package
from app.utils.validators import validate_egyptian_phone, validate_national_id

router = APIRouter(prefix="/import", tags=["استيراد Excel"])


@router.get("/template")
def download_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Customers"
    ws.append([
        "Customer Name", "National ID", "Phone", "Email", "Address",
        "Company Name", "Tax ID", "Package", "Subscription Type",
        "Contract Start", "Contract End", "Notes",
    ])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=customer_import_template.xlsx"},
    )


@router.post("/customers/validate")
async def validate_import(db: DbSession, user: CurrentUser, file: UploadFile = File(...)):
    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "يجب رفع ملف Excel")
    content = await file.read()
    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    job = ImportJob(job_type="customers", status="validating", total_rows=len(rows), created_by=user.id)
    db.add(job)
    db.flush()
    valid = invalid = duplicate = 0
    results = []
    for i, row in enumerate(rows, start=2):
        if not row or not row[0]:
            continue
        errors = []
        name, nid, phone = row[0], str(row[1] or ""), str(row[2] or "")
        ok_phone, err_phone = validate_egyptian_phone(phone)
        if not ok_phone:
            errors.append(err_phone)
        ok_nid, err_nid = validate_national_id(nid)
        if not ok_nid:
            errors.append(err_nid)
        is_dup = bool(db.scalar(select(Customer).where(Customer.national_id == nid, Customer.deleted_at.is_(None))))
        if is_dup:
            errors.append("عميل مكرر")
            duplicate += 1
        is_valid = len(errors) == 0 and not is_dup
        if is_valid:
            valid += 1
        else:
            invalid += 1
        import_row = ImportRow(
            job_id=job.id, row_number=i,
            data={"name": name, "national_id": nid, "phone": phone, "email": row[3], "address": row[4]},
            is_valid=is_valid, is_duplicate=is_dup, errors=errors,
        )
        db.add(import_row)
        results.append({"row": i, "valid": is_valid, "errors": errors})
    job.valid_rows = valid
    job.invalid_rows = invalid
    job.duplicate_rows = duplicate
    job.status = "validated"
    db.commit()
    return {
        "job_id": str(job.id),
        "total": len(rows),
        "valid": valid,
        "invalid": invalid,
        "duplicate": duplicate,
        "rows": results,
    }


@router.post("/customers/{job_id}/commit")
def commit_import(job_id: uuid.UUID, db: DbSession, user: CurrentUser):
    job = db.get(ImportJob, job_id)
    if not job:
        raise HTTPException(404)
    rows = db.scalars(select(ImportRow).where(ImportRow.job_id == job_id, ImportRow.is_valid.is_(True), ImportRow.imported.is_(False))).all()
    imported = 0
    for row in rows:
        data = row.data or {}
        count = db.scalar(select(Customer).where()) or 0
        code = f"C{date.today().year}{str(imported + 1).zfill(5)}"
        customer = Customer(
            customer_code=code,
            full_name=data.get("name", ""),
            national_id=data.get("national_id", ""),
            phone=data.get("phone", ""),
            email=data.get("email"),
            address=data.get("address"),
            status="active",
        )
        db.add(customer)
        row.imported = True
        imported += 1
    job.imported_rows = imported
    job.status = "completed"
    db.commit()
    return {"imported": imported}
