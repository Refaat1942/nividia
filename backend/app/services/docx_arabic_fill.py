import re
from copy import deepcopy

from docx import Document

# Replace Arabic contract blanks (dots/dashes) with values from context.
FILL_RULES: list[tuple[re.Pattern, str]] = [
    (re.compile(r"(الأستاذ\s*[:/]\s*)[.\s_\-]{3,}"), "customer_name"),
    (re.compile(r"(أقر\s+أنا\s*/\s*)[.\s_\-]{3,}"), "customer_name"),
    (re.compile(r"(رقم\s*قومي\s*)[.\s_\-]{3,}"), "national_id"),
    (re.compile(r"(العنوان\s*[:/]\s*)[.\s_\-]{3,}"), "address"),
    (re.compile(r"(صاحب\s+شركه\s*/\s*)[.\s_\-]{3,}"), "company_name"),
    (re.compile(r"(ممثل\s+شركه\s*)[.\s_\-]{0,30}"), "company_name"),
    (re.compile(r"مكتب\s*\(\s*\)"), "office_name_paren"),
    (re.compile(r"(استخدام\s+عدد\s+ساعات\s+سنوية\s*\()\s*[.\s_\-]{1,20}(\))"), "included_hours"),
    (re.compile(r"(بواقع\s*\()\s*[.\s_\-]{1,20}(\)\s*ساعات\s*شهرية)"), "monthly_hours"),
    (re.compile(r"(\d+\s*جنيه\s*فقط\s*\(\s*)[.\s_\-]{3,}(\s*جنيه\s*لا\s*غير)"), "package_price_words"),
    (re.compile(r"(القيمة\s+الإيجارية[^.]{0,40}هي\s*)[.\s_\-]{3,}(\s*جنية)"), "package_price"),
    (re.compile(r"(تبدأ\s+من\s*)[.\s_/]{3,}(\s*/\s*2026)"), "contract_start_ar"),
    (re.compile(r"(وتنتهي\s+في\s*)[.\s_/]{3,}(\s*/\s*2027)"), "contract_end_ar"),
    (re.compile(r"(في\s+يوم\s*)[.\s_\-]{3,}(\s*الموافق)"), "today_day"),
]


def _ctx_value(context: dict, key: str) -> str:
    if key == "office_name_paren":
        name = context.get("office_name") or context.get("room_name") or ""
        return f"مكتب ({name})" if name else "مكتب (—)"
    if key == "monthly_hours":
        try:
            annual = float(context.get("included_hours") or 0)
            return str(int(round(annual / 12))) if annual else "—"
        except (TypeError, ValueError):
            return "—"
    if key == "package_price_words":
        return str(context.get("package_price") or "—")
    if key == "contract_start_ar":
        d = str(context.get("contract_start") or "")
        parts = d.split("-")
        if len(parts) == 3:
            return f"{parts[2]} / {parts[1]} / {parts[0]}"
        return d
    if key == "contract_end_ar":
        d = str(context.get("contract_end") or "")
        parts = d.split("-")
        if len(parts) == 3:
            return f"{parts[2]} / {parts[1]} / {parts[0]}"
        return d
    if key == "today_day":
        from datetime import date
        return str(date.today().day)
    return str(context.get(key) or "—")


def _replace_in_paragraph(paragraph, context: dict) -> None:
    text = paragraph.text
    if not text.strip():
        return
    new_text = text
    for pattern, key in FILL_RULES:
        val = _ctx_value(context, key)
        if key in ("included_hours", "monthly_hours", "package_price", "package_price_words"):
            new_text = pattern.sub(lambda m, v=val: m.group(1) + v + (m.group(2) if m.lastindex and m.lastindex >= 2 else ""), new_text, count=1)
        elif key == "office_name_paren":
            new_text = pattern.sub(val, new_text, count=1)
        else:
            new_text = pattern.sub(lambda m, v=val: m.group(1) + v, new_text, count=1)
    if new_text != text:
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = new_text
        else:
            paragraph.add_run(new_text)


def _walk_paragraphs(doc: Document, context: dict) -> None:
    for para in doc.paragraphs:
        _replace_in_paragraph(para, context)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, context)


def fill_arabic_blanks(source_path: str, dest_path: str, context: dict) -> str:
    doc = Document(source_path)
    ctx = deepcopy(context)
    _walk_paragraphs(doc, ctx)
    doc.save(dest_path)
    return dest_path
