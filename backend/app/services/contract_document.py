"""Turn a DOCX contract into an editable, printable HTML document.

Every blank in the template (dots, dashes, underscores, empty brackets) becomes a
``<span data-fill>`` element.  The frontend lists those spans as form inputs so
staff can fill them, while the rest of the text stays freely editable.
"""

import html
import os
import re
from datetime import date, datetime
from html.parser import HTMLParser

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

# A blank is a run of dots / dashes / underscores (spaces allowed between them),
# or an empty pair of brackets.  Arabic text in the template sometimes uses the
# mirrored bracket direction, so both orders are accepted.
BLANK_RE = re.compile(
    r"(?=[.…]{2})[.…]+(?:[ \t]+[.…]+)*"
    r"|-{3,}"
    r"|_{3,}"
    r"|\([ \t]*\)"
    r"|\)[ \t]*\("
)
# Wide runs of spaces used as blanks after these words (e.g. "من      الى      ").
SPACE_BLANK_RE = re.compile(r"(?<=\sمن)[ \t]{3,}|(?<=\sالى)[ \t]{3,}")
# Places where the template expects a value but leaves only layout padding.
INSERT_AFTER_RE = re.compile(r"ممثل\s+شركه")

WEEKDAYS_AR = ["الاثنين", "الثلاثاء", "الأربعاء", "الخميس", "الجمعة", "السبت", "الأحد"]

# (regex on text just before the blank, key, Arabic label).  First match wins.
CONTEXT_RULES: list[tuple[re.Pattern, str, str]] = [
    (re.compile(r"في\s+يوم\s*$"), "day_name", "يوم التحرير"),
    (re.compile(r"الأستاذ\s*[:/]?\s*$"), "customer_name", "اسم الطرف الثاني"),
    (re.compile(r"رقم\s*قومي\s*$"), "national_id", "الرقم القومي"),
    (re.compile(r"العنوان\s*[:/]?\s*$"), "address", "العنوان"),
    (re.compile(r"ممثل\s+شركه\s*$"), "company_name", "اسم الشركة (ممثل شركة)"),
    (re.compile(r"(?:اقر|أقر)\s+أنا\s*/\s*$"), "customer_name", "اسم المُقر"),
    (re.compile(r"صاحب\s+شركه\s*/\s*$"), "company_name", "اسم الشركة (صاحب شركة)"),
    (re.compile(r"مكتب\s*$"), "office_name", "رقم / اسم المكتب"),
    (re.compile(r"هي\s*$"), "package_price", "القيمة الإيجارية (رقمًا)"),
    (re.compile(r"فقط\s*\(\s*$"), "package_price_words", "القيمة الإيجارية (كتابةً)"),
    (re.compile(r"ساعات\s+سنوية\s*\(\s*$"), "included_hours", "عدد الساعات السنوية"),
    (re.compile(r"بواقع\s*\(\s*$"), "monthly_hours", "عدد الساعات الشهرية"),
    (re.compile(r"الموافق\s+يوم\s*$"), "usage_day", "يوم الاستخدام"),
    (re.compile(r"\sمن\s*$"), "usage_from", "من الساعة"),
    (re.compile(r"\sالى\s*$"), "usage_to", "إلى الساعة"),
]

# Dates written as ".... / ... / 2026": (word before the date, context date key, label).
DATE_RE = re.compile(r"(?P<d>(?=[.…]{2})[.…]+(?:[ \t]+[.…]+)*)[ \t]*/[ \t]*(?P<m>(?=[.…]{2})[.…]+(?:[ \t]+[.…]+)*)[ \t]*/")
DATE_CONTEXT: list[tuple[re.Pattern, str, str]] = [
    (re.compile(r"الموافق\s*$"), "today_date", "تاريخ التحرير"),
    (re.compile(r"تبدأ\s+من\s*$"), "contract_start", "تاريخ بداية العقد"),
    (re.compile(r"(?:ينتهي|وينتهي)\s+في\s*$"), "contract_end", "تاريخ نهاية العقد"),
]


def _parse_date(value) -> date | None:
    if isinstance(value, date):
        return value
    try:
        return datetime.strptime(str(value or "")[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


def blank_value(key: str, context: dict) -> str:
    """Pre-fill value for a blank, or '' when the platform has no data for it."""
    ctx = context or {}
    if key == "day_name":
        d = _parse_date(ctx.get("today_date")) or date.today()
        return WEEKDAYS_AR[d.weekday()]
    if key.endswith(("__day", "__month")):
        base, part = key.rsplit("__", 1)
        d = _parse_date(ctx.get(base))
        if not d:
            return ""
        return f"{d.day:02d}" if part == "day" else f"{d.month:02d}"
    if key == "office_name":
        return str(ctx.get("office_name") or ctx.get("room_name") or "")
    if key == "monthly_hours":
        try:
            annual = float(ctx.get("included_hours") or 0)
        except (TypeError, ValueError):
            return ""
        return str(int(round(annual / 12))) if annual else ""
    if key == "included_hours":
        value = str(ctx.get("included_hours") or "")
        return "" if value in ("0", "0.0") else value
    return str(ctx.get(key) or "")


def _find_blanks(text: str) -> list[dict]:
    """Return blanks in a paragraph as dicts with start, end, key and label."""
    blanks: list[dict] = []
    taken: list[tuple[int, int]] = []

    def free(s: int, e: int) -> bool:
        return all(e <= a or s >= b for a, b in taken)

    for m in DATE_RE.finditer(text):
        before = text[: m.start()]
        for pattern, key, label in DATE_CONTEXT:
            if pattern.search(before):
                break
        else:
            key, label = "date", "تاريخ"
        for part, part_label in (("d", "اليوم"), ("m", "الشهر")):
            s, e = m.span(part)
            blanks.append({"start": s, "end": e, "key": f"{key}__{'day' if part == 'd' else 'month'}",
                           "label": f"{label} — {part_label}"})
            taken.append((s, e))

    candidates = [(m.start(), m.end()) for m in BLANK_RE.finditer(text)]
    candidates += [(m.start(), m.end()) for m in SPACE_BLANK_RE.finditer(text)]
    candidates += [(m.end(), m.end()) for m in INSERT_AFTER_RE.finditer(text)]
    for s, e in sorted(candidates):
        if not free(s, e) or (s == e and any(a <= s <= b for a, b in taken)):
            continue
        before = text[:s]
        key, label = "", ""
        for pattern, k, lbl in CONTEXT_RULES:
            if pattern.search(before):
                key, label = k, lbl
                break
        if not label:
            words = re.sub(r"[.…\-_()/:]+", " ", before).split()
            label = " ".join(words[-3:]) or "حقل"
        blanks.append({"start": s, "end": e, "key": key, "label": label})
        taken.append((s, e))
    blanks.sort(key=lambda b: (b["start"], b["end"]))
    return blanks


def _run_style(run) -> str:
    styles = []
    if run.bold:
        styles.append("font-weight:700")
    if run.italic:
        styles.append("font-style:italic")
    if run.underline:
        styles.append("text-decoration:underline")
    if run.font.size:
        styles.append(f"font-size:{run.font.size.pt:g}pt")
    return ";".join(styles)


# Long runs of spaces are layout padding in Word (e.g. "الطرف الأول      الطرف الثاني").
# Render them as a capped fixed-width gap so they do not wrap onto extra lines.
_GAP_RE = re.compile(r" {8,}")


def _gap(m: re.Match) -> str:
    return f'<span class="gap" style="width:{len(m.group(0)) * 0.25:g}em"> </span>'


def _paragraph_html(paragraph, context: dict, counter: list[int]) -> str:
    segments = [(r.text, _run_style(r)) for r in paragraph.runs if r.text]
    text = "".join(t for t, _ in segments)

    styles = []
    align = paragraph.alignment
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        styles.append("text-align:center")
    elif align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        styles.append("text-align:justify")
    if paragraph.paragraph_format.page_break_before:
        styles.append("break-before:page")
    style_attr = f' style="{";".join(styles)}"' if styles else ""

    if not text.strip():
        return f"<p{style_attr}><br></p>"

    # Character-level style lookup so blanks spanning several runs keep formatting.
    char_styles: list[str] = []
    for t, st in segments:
        char_styles.extend([st] * len(t))

    def emit(s: int, e: int) -> str:
        out, i = [], s
        while i < e:
            j = i
            while j < e and char_styles[j] == char_styles[i]:
                j += 1
            chunk = _GAP_RE.sub(_gap, html.escape(text[i:j], quote=False))
            out.append(f'<span style="{char_styles[i]}">{chunk}</span>' if char_styles[i] else chunk)
            i = j
        return "".join(out)

    parts, pos = [], 0
    for blank in _find_blanks(text):
        parts.append(emit(pos, blank["start"]))
        counter[0] += 1
        value = blank_value(blank["key"], context) if blank["key"] else ""
        st = char_styles[min(blank["start"], len(char_styles) - 1)] if char_styles else ""
        pad = " " if blank["start"] == blank["end"] else ""
        matched = text[blank["start"]:blank["end"]]
        open_br, close_br = (matched[0], matched[-1]) if matched[:1] in "()" and matched else ("", "")
        attrs = (
            f'data-fill="{counter[0]}" data-key="{html.escape(blank["key"])}" '
            f'data-label="{html.escape(blank["label"])}" title="{html.escape(blank["label"])}"'
        )
        if st:
            attrs += f' style="{st}"'
        parts.append(f"{pad}{open_br}<span {attrs}>{html.escape(value, quote=False)}</span>{close_br}{pad}")
        pos = blank["end"]
    parts.append(emit(pos, len(text)))
    return f"<p{style_attr}>{''.join(parts)}</p>"


def docx_to_editable_html(docx_path: str, context: dict | None = None) -> str:
    doc = Document(docx_path)
    counter = [0]
    out: list[str] = []
    body = doc.element.body
    paragraphs = {p._p: p for p in doc.paragraphs}
    tables = {t._tbl: t for t in doc.tables}
    for child in body.iterchildren():
        if child in paragraphs:
            out.append(_paragraph_html(paragraphs[child], context or {}, counter))
        elif child in tables:
            rows = []
            for row in tables[child].rows:
                cells = "".join(
                    "<td>" + "".join(_paragraph_html(p, context or {}, counter) for p in cell.paragraphs) + "</td>"
                    for cell in row.cells
                )
                rows.append(f"<tr>{cells}</tr>")
            out.append(f"<table>{''.join(rows)}</table>")
    return "".join(out)


_DANGEROUS_BLOCK_RE = re.compile(
    r"<(script|style|iframe|object|embed|link|meta|base|form|svg|math)\b[^>]*>.*?</\1\s*>"
    r"|<(script|style|iframe|object|embed|link|meta|base|form|input|svg|math)\b[^>]*/?>",
    re.IGNORECASE | re.DOTALL,
)
_EVENT_ATTR_RE = re.compile(r"\s+on[a-z]+\s*=\s*(\"[^\"]*\"|'[^']*'|[^\s>]+)", re.IGNORECASE)
_JS_URL_RE = re.compile(r"(href|src)\s*=\s*([\"']?)\s*(javascript|vbscript|data):", re.IGNORECASE)


def sanitize_document_html(raw: str) -> str:
    """Strip scripts, event handlers and script URLs from user-edited HTML."""
    cleaned = raw or ""
    prev = None
    while prev != cleaned:
        prev = cleaned
        cleaned = _DANGEROUS_BLOCK_RE.sub("", cleaned)
    cleaned = _EVENT_ATTR_RE.sub("", cleaned)
    cleaned = _JS_URL_RE.sub(r"\1=\2#", cleaned)
    return cleaned


# Shared by the in-app editor (frontend copy of these rules) and the print page.
DOCUMENT_CSS = """
.contract-doc { font-family: Arial, 'Segoe UI', Tahoma, sans-serif; font-size: 16pt; line-height: 1.25;
  color: #000; direction: rtl; text-align: right; white-space: pre-wrap; word-wrap: break-word; }
.contract-doc p { margin: 0 0 4pt 0; }
.contract-doc .gap { display: inline-block; max-width: 50%; }
.contract-doc table { width: 100%; border-collapse: collapse; }
.contract-doc td { border: 1px solid #000; padding: 4pt; vertical-align: top; }
.contract-doc [data-fill] { padding: 0 2px; }
.contract-doc [data-fill]:empty::before { content: '..............'; color: #000; }
.contract-doc .page-break { break-after: page; page-break-after: always; height: 0; }
"""


def print_page_html(title: str, body_html: str) -> str:
    return f"""<!DOCTYPE html>
<html dir="rtl" lang="ar">
<head>
  <meta charset="utf-8">
  <title>{html.escape(title)}</title>
  <style>
    @page {{ size: A4; margin: 2cm 2cm; }}
    body {{ margin: 0; background: #e2e8f0; }}
    .sheet {{ background: #fff; max-width: 21cm; margin: 24px auto; padding: 2cm; box-shadow: 0 2px 12px rgba(0,0,0,.15); box-sizing: border-box; }}
    {DOCUMENT_CSS}
    .actions {{ text-align: center; margin: 24px; font-family: Arial, sans-serif; }}
    .actions button {{ background: #1e40af; color: #fff; border: none; padding: 12px 24px; border-radius: 8px; cursor: pointer; font-size: 16px; }}
    @media print {{
      body {{ background: #fff; }}
      .sheet {{ margin: 0; padding: 0; box-shadow: none; max-width: none; }}
      .actions {{ display: none; }}
    }}
  </style>
</head>
<body>
  <div class="actions"><button onclick="window.print()">طباعة</button></div>
  <div class="sheet"><div class="contract-doc">{body_html}</div></div>
  <script>window.addEventListener('load', function () {{ setTimeout(function () {{ window.print(); }}, 300); }});</script>
</body>
</html>"""


# --- Edited HTML -> DOCX ------------------------------------------------------

_BLOCK_TAGS = {"p", "div", "h1", "h2", "h3", "h4", "li", "tr"}
_EMPTY_FILL = ".............."


class _HtmlToBlocks(HTMLParser):
    """Flatten editor HTML into paragraphs of styled text runs."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.blocks: list[dict] = []
        self.current: dict | None = None
        self.stack: list[tuple[str, dict, bool]] = []  # (tag, style, is_empty_fill)
        self.fill_had_text: list[bool] = []

    def _style(self) -> dict:
        merged: dict = {}
        for _, st, _ in self.stack:
            merged.update(st)
        return merged

    def _ensure_block(self) -> dict:
        if self.current is None:
            self.current = {"runs": [], "align": None, "page_break_before": False}
            self.blocks.append(self.current)
        return self.current

    def _close_block(self) -> None:
        self.current = None

    def handle_starttag(self, tag, attrs):
        a = dict(attrs)
        css = _parse_css(a.get("style") or "")
        cls = a.get("class") or ""
        if tag in _BLOCK_TAGS:
            if "page-break" in cls:
                self._close_block()
                self.blocks.append({"page_break": True})
                return
            self._close_block()
            block = self._ensure_block()
            align = css.get("text-align") or a.get("align")
            if align:
                block["align"] = align
            return
        if tag == "br":
            self._ensure_block()["runs"].append({"text": "\n", "style": self._style()})
            return
        style: dict = {}
        if tag in ("b", "strong") or css.get("font-weight") in ("700", "bold", "800", "900"):
            style["bold"] = True
        if css.get("font-weight") in ("400", "normal"):
            style["bold"] = False
        if tag in ("u", "ins") or "underline" in css.get("text-decoration", ""):
            style["underline"] = True
        if tag in ("i", "em") or css.get("font-style") == "italic":
            style["italic"] = True
        size = css.get("font-size", "")
        m = re.match(r"([\d.]+)pt", size)
        if m:
            style["size"] = float(m.group(1))
        if "gap" in cls.split():
            w = re.match(r"([\d.]+)em", css.get("width", ""))
            spaces = int(round(float(w.group(1)) / 0.25)) if w else 8
            self._ensure_block()["runs"].append({"text": " " * spaces, "style": self._style()})
            self.stack.append((tag, style, False))
            self.stack.append(("__gap__", {}, False))
            return
        is_fill = "data-fill" in a
        self.stack.append((tag, style, is_fill))
        if is_fill:
            self.fill_had_text.append(False)

    def handle_endtag(self, tag):
        if tag in _BLOCK_TAGS:
            self._close_block()
            return
        if self.stack and self.stack[-1][0] == "__gap__":
            self.stack.pop()
        for i in range(len(self.stack) - 1, -1, -1):
            if self.stack[i][0] == tag:
                _, _, is_fill = self.stack[i]
                if is_fill and self.fill_had_text:
                    had = self.fill_had_text.pop()
                    if not had:
                        self._ensure_block()["runs"].append({"text": _EMPTY_FILL, "style": self._style()})
                del self.stack[i:]
                break

    def handle_data(self, data):
        if not data:
            return
        if self.stack and self.stack[-1][0] == "__gap__":
            return
        if self.current is None and not data.strip():
            return  # whitespace between blocks
        if self.fill_had_text and data.strip():
            self.fill_had_text[-1] = True
        self._ensure_block()["runs"].append({"text": data, "style": self._style()})


def _parse_css(style: str) -> dict:
    out = {}
    for decl in style.split(";"):
        if ":" in decl:
            k, v = decl.split(":", 1)
            out[k.strip().lower()] = v.strip().lower()
    return out


def _set_rtl_run(run, style: dict) -> None:
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attr), "Arial")
    if style.get("bold"):
        run.bold = True
        rpr.append(OxmlElement("w:bCs"))
    if style.get("italic"):
        run.italic = True
    if style.get("underline"):
        run.underline = True
    size = style.get("size") or 16
    run.font.size = Pt(size)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(round(size * 2))))
    rpr.append(sz_cs)
    rpr.append(OxmlElement("w:rtl"))


def html_to_docx(template_path: str, document_html: str, output_path: str) -> str:
    """Write the edited contract HTML into a DOCX that keeps the template's page setup."""
    parser = _HtmlToBlocks()
    parser.feed(document_html or "")
    parser.close()

    doc = Document(template_path)
    body = doc.element.body
    for child in list(body.iterchildren()):
        if child.tag != qn("w:sectPr"):
            body.remove(child)

    pending_break = False
    for block in parser.blocks:
        if block.get("page_break"):
            pending_break = True
            continue
        para = doc.add_paragraph()
        ppr = para._p.get_or_add_pPr()
        ppr.append(OxmlElement("w:bidi"))
        para.paragraph_format.space_after = Pt(4)
        if pending_break:
            para.paragraph_format.page_break_before = True
            pending_break = False
        align = block.get("align")
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "justify":
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif align == "left":
            # In a bidi paragraph Word's "right" is the visual left edge.
            para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        runs = block["runs"]
        if len(runs) == 1 and runs[0]["text"] == "\n":
            runs = []  # the editor's "<p><br></p>" is just an empty paragraph
        for run_data in runs:
            pieces = run_data["text"].split("\n")
            for idx, piece in enumerate(pieces):
                if idx:
                    para.add_run().add_break()
                if piece:
                    _set_rtl_run(para.add_run(piece), run_data["style"])

    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    doc.save(output_path)
    return output_path
